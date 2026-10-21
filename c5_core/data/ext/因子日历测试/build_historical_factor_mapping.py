from __future__ import annotations

import csv
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


BASE = Path(__file__).resolve().parent
OUTPUT_DIR = BASE / "historical_calendar_backtest_results"


# 名称不同但公式相同的项目。这里仅合并计算字段、方向、窗口定义一致的公式。
FORMULA_ALIASES = [
    ["Amihud 非流动性", "Amihud 非流动性因子"],
    ["K 线最短路径非流动性", "非流动性（K 线最短路径）"],
    ["一致买入交易", "一致买入交易 PCV", "一致买入交易因子 PCV"],
    ["一致交易", "一致交易 TCV"],
    ["一致卖出交易", "一致卖出交易 NCV"],
    ["下行 Beta", "下行贝塔"],
    ["主力交易强度", "主力交易强度 MTS"],
    ["主力交易情绪", "主力交易情绪 MTE"],
    ["乖离率", "乖离率 BIAS"],
    ["价量相关性因子", "价量相关性因子 CPV"],
    ["价量相关性平均数因子", "价量相关性平均数因子 PV_corr_avg"],
    ["价量相关性波动性因子", "价量相关性波动性因子 PV_corr_std"],
    ["价量相关性综合因子", "价量相关性综合因子 PV_corr"],
    ["价量相关性趋势因子", "价量相关性趋势因子 PV_corr_trend"],
    ["估波指标", "估波指标 Coppock"],
    ["佳庆指标", "佳庆指标 CHO"],
    ["佳庆离散指标", "佳庆离散指标 CV", "蔡金波动率指标 CVI"],
    ["信息离散度", "信息离散度 ID"],
    ["克林格成交量摆动指标 KVO", "克林格成交量摆动指标 KVO（重复页合并）"],
    ["动态买卖气指标", "动态买卖气指标 ADTM"],
    ["动态动量指标 DMI", "相对强弱指标（原文 DMI 公式）"],
    ["单笔成交金额分位数", "单笔成交金额分位数 QUA"],
    ["变动速率", "变动速率 ROC"],
    ["在险价值", "在险价值 VaR"],
    ["基于日度收益的市场 Beta", "市场 Beta"],
    ["多空指数", "多空指数 BBI"],
    ["对数流通市值", "对数自由流通市值"],
    ["平均真实波幅", "平均真实波幅 ATR"],
    ["异同离差乖离率", "异同离差乖离率 DBCD"],
    ["心理线", "心理线 PSY"],
    ["基于振幅切割的动量", "振幅切割后的长端动量"],
    ["换手率分布均匀度 UTD", "换手率分布均匀度因子"],
    ["换手率调整的日成交量为零的数量", "换手率调整的日成交量为零的数量 Lm"],
    ["时间加权平均的相对价格位置 ARPP", "时间加权平均的股票相对价格位置", "时间加权平均的股票相对价格位置 ARPP"],
    ["时间序列动量", "时间序列动量 TSMOM"],
    ["最近 52 周最高价因子", "最近 52 周的最高价"],
    ["梅斯线", "梅斯线 Mass Index"],
    ["流动性冲击", "流动性冲击 LIQU"],
    ["流通股市值", "自由流通股市值"],
    ["理想反转", "理想反转因子"],
    ["累积/派发线", "累积/派发线 AD"],
    ["综合股权发行-年度", "综合股权发行（年度）"],
    ["负偏度系数", "负偏度系数 NCSKEW"],
    ["资本利得突出量", "资本收益过剩 CGO"],
    ["阿隆指数 AROON", "阿隆指标 AROON"],
    ["随机指标", "随机指标 KDJ"],
    ["非流动性的变异系数", "非流动性的变异系数 CVILLIQ"],
    ["散户买卖非平衡性 BSI", "散户的买卖非平衡性 BSI"],
    ["成交额变动", "成交额波动"],
    ["高频已实现偏度", "高频收益偏度"],
]


# 名称不同、但可直接使用现有代码的公式。
DIRECT_ALIASES = {
    "K 线最短路径非流动性": "ILLIQ_SHORTCUT",
    "非流动性（K 线最短路径）": "ILLIQ_SHORTCUT",
    "“午蔽古木”因子": "A15_NOON_ANCIENT_TREE",
    "“孤雁出群”因子": "F02_D12_LONE_GOOSE",
    "“情绪溢出”因子": "J09_D04_EMOTION_SPILLOVER",
    "高频下行波动占比": "A21_DOWNSIDE_VOL_SHARE_5M_20D",
    "高频已实现峰度": "F02_D16_REALIZED_KURTOSIS_1M_20D",
    "趋势占比": "J09_D10_INTRADAY_TREND_RATIO",
    "置信正态分布主动占比": "J11_D11_NORMAL_CDF_ACTIVE_SHARE",
    "聪明钱因子": "J09_D14_SMART_MONEY_VWAP_RATIO",
    "隔夜收益率": "J11_D17_OVERNIGHT_RETURN",
    "已实现双幂次变差": "J09_D17_REALIZED_BIPOWER_VARIATION",
    "成交量“潮汐”的价格变动速率": "J11_D14_VOLUME_TIDE_PRICE_RATE",
    "成交量波峰计数因子": "J11_D01_VOLUME_PEAK_COUNT",
    "持续异常交易量": "J06_D13_PATV_5M",
    "朴素主动占比": "J09_D11_NAIVE_ACTIVE_SHARE",
    "一致买入交易": "J11_D20_CONSISTENT_BUYING",
    "一致买入交易 PCV": "J11_D20_CONSISTENT_BUYING",
    "一致买入交易因子 PCV": "J11_D20_CONSISTENT_BUYING",
    "大的上行跳跃波动率": "M16_LARGE_UP_JUMP_VOL",
}


ORDINARY_UNAVAILABLE = {
    "Dimson 贝塔",
    "协偏度（系统偏度）",
    "尾部 Beta",
    "月度特质波动率",
    "残差动量-基于 CAPM",
    "残差动量-基于 FF3",
    "特异度",
    "特质偏度",
    "特质波动率",
    "行业动量-横向切割",
    "行业动量-纵向切割",
    "盈余公告回报",
    "盈余公告股票跳动",
}


def read_source_rows() -> list[dict[str, str]]:
    candidates = list((BASE / "因子日历Markdown" / "因子清单").glob("*.md"))
    source = max(candidates, key=lambda path: path.stat().st_size)
    rows: list[dict[str, str]] = []
    for line in source.read_text(encoding="utf-8").splitlines():
        if not line.startswith("| 202"):
            continue
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if len(cells) != 9:
            continue
        rows.append(
            {
                "year": cells[0],
                "date": cells[1],
                "name": cells[2],
                "formula_text": cells[3],
                "required_fields": cells[4],
                "source_category": cells[5],
                "backtest": cells[6],
                "reason": cells[7],
                "source": cells[8],
            }
        )
    if len(rows) != 524:
        raise RuntimeError(f"应读取524条记录，实际读取{len(rows)}条。")
    return rows


def normalize_name(value: str) -> str:
    value = re.sub(r"[（(][^）)]*[）)]", "", value)
    value = "".join(
        char for char in value if "\u4e00" <= char <= "\u9fff" or char.isalnum()
    )
    for word in ("因子", "指标", "原文", "解释版本", "分钟", "高频"):
        value = value.replace(word, "")
    return value.lower()


def read_current_catalog() -> tuple[dict[str, list[dict[str, str]]], dict[str, dict[str, str]]]:
    workbook = next((BASE / "outputs" / "minute_completion").glob("*.xlsx"))
    sheet = load_workbook(workbook, read_only=False, data_only=True).worksheets[1]
    by_name: dict[str, list[dict[str, str]]] = defaultdict(list)
    by_code: dict[str, dict[str, str]] = {}
    for row in sheet.iter_rows(min_row=5, values_only=True):
        if row[0] is None:
            continue
        item = {
            "name": str(row[1]),
            "code": str(row[2]),
            "signal_frequency": "" if row[6] is None else str(row[6]),
            "status": "" if row[14] is None else str(row[14]),
        }
        by_name[normalize_name(item["name"])].append(item)
        by_code[item["code"]] = item
    return by_name, by_code


def make_formula_keys(rows: list[dict[str, str]]) -> dict[int, str]:
    parent = list(range(len(rows)))

    def find(index: int) -> int:
        while parent[index] != index:
            parent[index] = parent[parent[index]]
            index = parent[index]
        return index

    def union(left: int, right: int) -> None:
        left_root = find(left)
        right_root = find(right)
        if left_root != right_root:
            parent[right_root] = left_root

    # 同名通常是跨年重复公式；两个已确认公式方向不同的例外单独保留。
    same_name: dict[str, list[int]] = defaultdict(list)
    for index, row in enumerate(rows):
        special = ""
        if row["name"] == "极端下行风险" and row["year"] == "2022":
            special = "_raw_return_evt"
        if row["name"] in {"残差反转", "残差反转因子"} and row["year"] == "2024":
            special = "_reverse_on_flow"
        same_name[row["name"] + special].append(index)
    for indices in same_name.values():
        for index in indices[1:]:
            union(indices[0], index)

    rows_by_name: dict[str, list[int]] = defaultdict(list)
    for index, row in enumerate(rows):
        rows_by_name[row["name"]].append(index)
    for aliases in FORMULA_ALIASES:
        indices = [index for name in aliases for index in rows_by_name.get(name, [])]
        # 2024残差反转的回归方向与另外两年不同，不参与别名合并。
        if "残差反转" in aliases or "残差反转因子" in aliases:
            indices = [index for index in indices if rows[index]["year"] != "2024"]
        if indices:
            for index in indices[1:]:
                union(indices[0], index)

    roots = sorted({find(index) for index in range(len(rows))})
    key_by_root = {root: f"HF{position:04d}" for position, root in enumerate(roots, 1)}
    return {index: key_by_root[find(index)] for index in range(len(rows))}


def infer_frequency(row: dict[str, str], current: dict[str, str] | None) -> str:
    if current and current["signal_frequency"]:
        return current["signal_frequency"]
    text = row["formula_text"] + row["required_fields"]
    if "月末" in text or "个月" in text or "月度" in text or "按月" in text:
        return "monthly"
    if "每周" in text or "周度" in text:
        return "weekly"
    return "daily"


def ordinary_missing(row: dict[str, str]) -> str:
    name = row["name"]
    if name == "极端下行风险" and row["year"] != "2022":
        return "缺无风险利率及Fama-French因子序列"
    if name in ORDINARY_UNAVAILABLE:
        if "行业" in name:
            return "缺历史行业分类或行业指数"
        if "盈余公告" in name:
            return "缺历史公告日期及披露时点"
        return "缺无风险利率或Fama-French因子序列"
    return ""


def classify_kind(row: dict[str, str], current: dict[str, str] | None) -> str:
    if ordinary_missing(row):
        return "unavailable"
    if current:
        return "existing"
    text = row["formula_text"] + row["required_fields"]
    risk_words = ("回归", "横截面", "市场收益", "股票池", "相关网络", "极值分布", "条件在险价值")
    if any(word in text for word in risk_words):
        return "risk"
    cache_words = ("分钟", "5 分钟", "10 分钟", "日内", "跳跃", "已实现", "成交量占比")
    if any(word in text for word in cache_words):
        return "cached"
    return "standard"


def current_match(
    row: dict[str, str],
    current_by_name: dict[str, list[dict[str, str]]],
    current_by_code: dict[str, dict[str, str]],
) -> dict[str, str] | None:
    alias_code = DIRECT_ALIASES.get(row["name"])
    if alias_code:
        return current_by_code.get(alias_code, {"code": alias_code, "signal_frequency": "", "status": ""})
    matches = current_by_name.get(normalize_name(row["name"]), [])
    return matches[0] if matches else None


def main() -> None:
    rows = read_source_rows()
    formula_keys = make_formula_keys(rows)
    current_by_name, current_by_code = read_current_catalog()
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    ordinary_fields = [
        "record_id",
        "year",
        "date",
        "name",
        "formula_key",
        "formula_text",
        "category",
        "mode",
        "implementation_kind",
        "implementation_ref",
        "signal_frequency",
        "note",
    ]
    excluded_fields = ordinary_fields + ["required_fields", "excluded_reason"]
    ordinary: list[dict[str, str]] = []
    excluded: list[dict[str, str]] = []

    for index, row in enumerate(rows, 1):
        current = current_match(row, current_by_name, current_by_code)
        common = {
            "record_id": f"HC{index:04d}",
            "year": row["year"],
            "date": row["date"],
            "name": row["name"],
            "formula_key": formula_keys[index - 1],
            "formula_text": row["formula_text"],
            "category": "level2" if "Level2" in row["source_category"] else "ordinary_minute",
            "implementation_ref": current["code"] if current else "",
            "signal_frequency": infer_frequency(row, current),
        }
        if "Level2" in row["source_category"]:
            excluded.append(
                {
                    **common,
                    "mode": "unavailable",
                    "implementation_kind": "unavailable",
                    "note": "按用户要求，Level2数据准备完成前不开展回测。",
                    "required_fields": row["required_fields"],
                    "excluded_reason": row["reason"],
                }
            )
            continue

        missing_note = ordinary_missing(row)
        current_status = current["status"] if current else ""
        trial = (
            "近似" in row["source_category"]
            or "待补参数" in current_status
            or "未回测" in current_status
            or "待回测" in current_status
            or "已试算" in current_status
        )
        note_parts = []
        if missing_note:
            note_parts.append(missing_note)
        if current_status:
            note_parts.append(f"2026状态：{current_status}")
        if not note_parts:
            note_parts.append("公式和所需字段已记录，参数以原文为准。")
        ordinary.append(
            {
                **common,
                "mode": "unavailable" if missing_note else ("trial" if trial else "strict"),
                "implementation_kind": classify_kind(row, current),
                "note": "；".join(note_parts),
            }
        )

    if len(ordinary) != 415 or len(excluded) != 109:
        raise RuntimeError(f"行数不正确：普通{len(ordinary)}，Level2 {len(excluded)}。")

    with (OUTPUT_DIR / "factor_mapping.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=ordinary_fields)
        writer.writeheader()
        writer.writerows(ordinary)
    with (OUTPUT_DIR / "level2_excluded.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=excluded_fields)
        writer.writeheader()
        writer.writerows(excluded)

    ordinary_formula_count = len({row["formula_key"] for row in ordinary})
    excluded_formula_count = len({row["formula_key"] for row in excluded})
    print(
        f"ordinary_rows={len(ordinary)} ordinary_formulas={ordinary_formula_count} "
        f"level2_rows={len(excluded)} level2_formulas={excluded_formula_count}"
    )


if __name__ == "__main__":
    main()
