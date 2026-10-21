# -*- coding: utf-8 -*-
"""整理历史三版与 2026 版需要 Level2 数据的因子清单。"""

from __future__ import annotations

import argparse
import csv
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Iterable, Sequence

from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parent
DEFAULT_HISTORICAL_SUMMARY = (
    BASE_DIR / "historical_calendar_backtest_results" / "factor_backtest_summary.csv"
)
DEFAULT_2026_WORKBOOK = (
    BASE_DIR
    / "outputs"
    / "minute_completion"
    / "因子日历2026因子库_分钟因子补算完成_130个.xlsx"
)
DEFAULT_OUTPUT = BASE_DIR / "combined_level2_factor_catalog.csv"

EXPECTED_HISTORICAL_ROWS = 109
EXPECTED_HISTORICAL_UNITS = 77
EXPECTED_2026_CODES = 69
EXPECTED_COMBINED_ROWS = 178

PENDING_DATA_STATUS = "待补数据"
SPECIAL_2026_NOTES = {
    "J12_D31_PRICE_IMPACT_BIAS": "公式印刷错误，暂不能严格计算。",
}

# 成交笔数虽不一定来自逐笔文件，但普通分钟行情无法提供，也纳入待补清单。
LEVEL2_REQUIRED_HINTS = (
    "逐笔",
    "委托",
    "盘口",
    "买一",
    "卖一",
    "五档",
    "集合竞价",
    "成交笔数",
    "成交单数",
    "订单",
    "主买",
    "主卖",
    "主动买",
    "主动卖",
    "主动净流入",
    "买卖方向",
    "BS方向",
    "撤单",
    "DealNum",
    "Level2",
    "每笔",
    "单笔",
)

CATALOG_COLUMNS = (
    "source_scope",
    "source_edition",
    "source_record_id",
    "year",
    "date",
    "name",
    "code_or_unit",
    "factor_code",
    "formula_key",
    "unit_id",
    "formula",
    "source_frequency",
    "signal_frequency",
    "direction",
    "required_data",
    "source_reference",
    "backtest_status",
    "availability_note",
    "catalog_note",
    "historical_implementation_ref",
    "historical_record_ids",
    "historical_formula_keys",
    "historical_unit_ids",
    "historical_match",
)

WORKBOOK_REQUIRED_COLUMNS = {
    "日期",
    "因子名称",
    "因子代码",
    "原始时间频度",
    "建议/本次信号频度",
    "文献方向",
    "公式/计算方法",
    "所需数据",
    "当前数据情况",
    "当前可回测性",
    "回测状态",
    "实现说明",
    "来源",
    "限制/备注",
}


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _join_unique(values: Iterable[str]) -> str:
    return ";".join(sorted({value for value in values if value}))


def read_historical_level2(path: Path = DEFAULT_HISTORICAL_SUMMARY) -> list[dict[str, str]]:
    """读取历史汇总表中的全部 Level2 原始记录。"""
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        required = {
            "record_id",
            "year",
            "date",
            "name",
            "formula_key",
            "formula_text",
            "category",
            "implementation_ref",
            "signal_frequency",
            "unit_id",
            "direction",
            "required_fields",
            "excluded_reason",
        }
        missing = required - set(reader.fieldnames or [])
        if missing:
            raise ValueError(f"历史汇总表缺少字段：{sorted(missing)}")
        return [dict(row) for row in reader if row["category"].strip() == "level2"]


def read_2026_factor_library(path: Path = DEFAULT_2026_WORKBOOK) -> list[dict[str, str]]:
    """读取 2026 工作簿的因子总库。"""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if "因子总库" not in workbook.sheetnames:
            raise ValueError("2026 工作簿缺少“因子总库”工作表。")
        sheet = workbook["因子总库"]
        header_values = next(sheet.iter_rows(min_row=4, max_row=4, values_only=True))
        headers = [_text(value) for value in header_values]
        missing = WORKBOOK_REQUIRED_COLUMNS - set(headers)
        if missing:
            raise ValueError(f"因子总库缺少字段：{sorted(missing)}")

        rows: list[dict[str, str]] = []
        for values in sheet.iter_rows(min_row=5, values_only=True):
            if not values or values[0] is None:
                continue
            rows.append(
                {
                    header: _text(values[index]) if index < len(values) else ""
                    for index, header in enumerate(headers)
                    if header
                }
            )
        return rows
    finally:
        workbook.close()


def requires_level2_data(row: dict[str, str]) -> bool:
    """判断待补项目是否确实依赖逐笔、委托、盘口等数据。"""
    fields = (
        "因子名称",
        "公式/计算方法",
        "所需数据",
        "当前数据情况",
        "当前可回测性",
        "实现说明",
        "限制/备注",
    )
    description = "|".join(row.get(field, "") for field in fields)
    return any(hint in description for hint in LEVEL2_REQUIRED_HINTS)


def select_2026_level2_targets(rows: Iterable[dict[str, str]]) -> list[dict[str, str]]:
    """选出 2026 版需要补 Level2 数据的因子。"""
    selected: list[dict[str, str]] = []
    for row in rows:
        code = row.get("因子代码", "")
        status = row.get("回测状态", "")
        if status != PENDING_DATA_STATUS and code not in SPECIAL_2026_NOTES:
            continue
        if requires_level2_data(row):
            selected.append(row)

    selected.sort(key=lambda row: (row.get("日期", ""), row.get("因子代码", "")))
    codes = [row.get("因子代码", "") for row in selected]
    if not all(codes):
        raise ValueError("2026 Level2 清单存在空因子代码。")
    if len(codes) != len(set(codes)):
        raise ValueError("2026 Level2 清单存在重复因子代码。")
    return selected


def validate_source_counts(
    historical_rows: Sequence[dict[str, str]],
    current_rows: Sequence[dict[str, str]],
) -> None:
    """检查当前正式来源的固定数量。"""
    historical_units = {row["unit_id"] for row in historical_rows if row["unit_id"]}
    current_codes = {row["因子代码"] for row in current_rows if row["因子代码"]}
    if len(historical_rows) != EXPECTED_HISTORICAL_ROWS:
        raise ValueError(
            f"历史 Level2 应为{EXPECTED_HISTORICAL_ROWS}条，实际为{len(historical_rows)}条。"
        )
    if len(historical_units) != EXPECTED_HISTORICAL_UNITS:
        raise ValueError(
            f"历史 Level2 应为{EXPECTED_HISTORICAL_UNITS}个计算单元，"
            f"实际为{len(historical_units)}个。"
        )
    if len(current_rows) != EXPECTED_2026_CODES or len(current_codes) != EXPECTED_2026_CODES:
        raise ValueError(
            f"2026 Level2 应为{EXPECTED_2026_CODES}个代码，实际为{len(current_codes)}个。"
        )


def _history_output_row(row: dict[str, str]) -> dict[str, str]:
    implementation_ref = row.get("implementation_ref", "")
    return {
        "source_scope": "historical",
        "source_edition": f"历史因子日历{row.get('year', '')}",
        "source_record_id": row.get("record_id", ""),
        "year": row.get("year", ""),
        "date": row.get("date", ""),
        "name": row.get("name", ""),
        "code_or_unit": row.get("unit_id", ""),
        "factor_code": implementation_ref,
        "formula_key": row.get("formula_key", ""),
        "unit_id": row.get("unit_id", ""),
        "formula": row.get("formula_text", ""),
        "source_frequency": "",
        "signal_frequency": row.get("signal_frequency", ""),
        "direction": row.get("direction", ""),
        "required_data": row.get("required_fields", ""),
        "source_reference": "historical_calendar_backtest_results/factor_backtest_summary.csv",
        "backtest_status": row.get("backtest_status", ""),
        "availability_note": row.get("excluded_reason", ""),
        "catalog_note": row.get("note", ""),
        "historical_implementation_ref": implementation_ref,
        "historical_record_ids": row.get("record_id", ""),
        "historical_formula_keys": row.get("formula_key", ""),
        "historical_unit_ids": row.get("unit_id", ""),
        "historical_match": "历史原记录",
    }


def _current_output_row(
    row: dict[str, str],
    historical_by_ref: dict[str, list[dict[str, str]]],
) -> dict[str, str]:
    code = row["因子代码"]
    matches = historical_by_ref.get(code, [])
    availability_parts = [row.get("当前数据情况", ""), row.get("当前可回测性", "")]
    availability_note = "；".join(part for part in availability_parts if part)
    catalog_note = row.get("限制/备注", "") or row.get("实现说明", "")
    special_note = SPECIAL_2026_NOTES.get(code)
    if special_note:
        catalog_note = special_note + (f"{catalog_note}" if catalog_note else "")

    return {
        "source_scope": "current_2026",
        "source_edition": "因子日历2026",
        "source_record_id": code,
        "year": row.get("日期", "")[:4],
        "date": row.get("日期", ""),
        "name": row.get("因子名称", ""),
        "code_or_unit": code,
        "factor_code": code,
        "formula_key": "",
        "unit_id": "",
        "formula": row.get("公式/计算方法", ""),
        "source_frequency": row.get("原始时间频度", ""),
        "signal_frequency": row.get("建议/本次信号频度", ""),
        "direction": row.get("文献方向", ""),
        "required_data": row.get("所需数据", ""),
        "source_reference": row.get("来源", ""),
        "backtest_status": row.get("回测状态", ""),
        "availability_note": availability_note,
        "catalog_note": catalog_note,
        "historical_implementation_ref": code if matches else "",
        "historical_record_ids": _join_unique(match.get("record_id", "") for match in matches),
        "historical_formula_keys": _join_unique(match.get("formula_key", "") for match in matches),
        "historical_unit_ids": _join_unique(match.get("unit_id", "") for match in matches),
        "historical_match": (
            "2026代码与历史implementation_ref相同"
            if matches
            else "历史表无同代码implementation_ref"
        ),
    }


def build_combined_catalog(
    historical_path: Path = DEFAULT_HISTORICAL_SUMMARY,
    workbook_path: Path = DEFAULT_2026_WORKBOOK,
    *,
    check_expected_counts: bool = True,
) -> list[dict[str, str]]:
    """生成历史 109 条加 2026 年 69 个代码的合并清单。"""
    historical_rows = read_historical_level2(historical_path)
    current_rows = select_2026_level2_targets(read_2026_factor_library(workbook_path))
    if check_expected_counts:
        validate_source_counts(historical_rows, current_rows)

    historical_by_ref: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in historical_rows:
        implementation_ref = row.get("implementation_ref", "")
        if implementation_ref:
            historical_by_ref[implementation_ref].append(row)

    combined = [_history_output_row(row) for row in historical_rows]
    combined.extend(_current_output_row(row, historical_by_ref) for row in current_rows)
    if check_expected_counts and len(combined) != EXPECTED_COMBINED_ROWS:
        raise ValueError(
            f"合并清单应为{EXPECTED_COMBINED_ROWS}条，实际为{len(combined)}条。"
        )
    return combined


def write_catalog(rows: Iterable[dict[str, str]], output_path: Path) -> None:
    """写出带 BOM 的 UTF-8 CSV，便于 Excel 直接打开。"""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=CATALOG_COLUMNS)
        writer.writeheader()
        writer.writerows(rows)


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--historical-summary",
        type=Path,
        default=DEFAULT_HISTORICAL_SUMMARY,
        help="历史回测汇总 CSV",
    )
    parser.add_argument(
        "--workbook",
        type=Path,
        default=DEFAULT_2026_WORKBOOK,
        help="2026 因子总库工作簿",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=DEFAULT_OUTPUT,
        help="合并清单输出路径",
    )
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_args(argv)
    rows = build_combined_catalog(args.historical_summary, args.workbook)
    write_catalog(rows, args.output)
    historical_rows = [row for row in rows if row["source_scope"] == "historical"]
    current_rows = [row for row in rows if row["source_scope"] == "current_2026"]
    historical_units = {row["unit_id"] for row in historical_rows}
    print(f"output={args.output.resolve()}")
    print(
        f"historical_records={len(historical_rows)} "
        f"historical_units={len(historical_units)} "
        f"current_2026_codes={len(current_rows)} combined_records={len(rows)}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
