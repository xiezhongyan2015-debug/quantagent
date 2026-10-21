from __future__ import annotations

import argparse
import re
from pathlib import Path

import openpyxl
import pandas as pd

from wind_excel_client import WindExcelClient, excel_dt


DEFAULT_DAILY = Path(r"C:\Users\Administrator\Desktop\daily.xlsx")
CODE_RE = re.compile(r"^\d{6}\.(?:SZ|SH|BJ)$", re.I)


def read_daily_codes(path: Path) -> list[str]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb[wb.sheetnames[0]]
        codes = []
        for row in ws.iter_rows(min_row=4, max_row=4, values_only=True):
            for value in row:
                if value is None:
                    continue
                code = str(value).strip().upper()
                if CODE_RE.match(code):
                    codes.append(code)
        return list(dict.fromkeys(codes))
    finally:
        wb.close()


def chunks(items: list[str], size: int):
    for i in range(0, len(items), size):
        yield items[i : i + size]


def build(daily_path: Path, output: Path, batch: int = 100) -> pd.DataFrame:
    codes = read_daily_codes(daily_path)
    if not codes:
        raise RuntimeError(f"No A-share codes found in {daily_path}")
    fields = [
        "s_info_name",
        "s_ipo_listeddate",
        "s_info_delistdate",
        "s_info_exchmarket",
        "s_info_industry_sw",
        "s_info_compname",
        "s_info_province",
        "s_info_business",
    ]
    client = WindExcelClient()
    parts: list[pd.DataFrame] = []
    try:
        for batch_no, batch_codes in enumerate(chunks(codes, batch), 1):
            try:
                df = client.query_wss(batch_codes, fields, timeout=180)
            except Exception as exc:
                print(f"batch {batch_no} failed, retrying as halves: {exc}", flush=True)
                mid = len(batch_codes) // 2
                frames = []
                for half in (batch_codes[:mid], batch_codes[mid:]):
                    if half:
                        frames.append(client.query_wss(half, fields, timeout=180))
                df = pd.concat(frames, ignore_index=True)
            df.insert(0, "symbol", batch_codes[: len(df)])
            parts.append(df)
            print(f"static batch {batch_no}: {len(batch_codes)} codes", flush=True)
    finally:
        client.close()
    all_df = pd.concat(parts, ignore_index=True)
    for col in ("s_ipo_listeddate", "s_info_delistdate"):
        all_df[col] = pd.to_datetime(all_df[col], errors="coerce").dt.date
    all_df = all_df.rename(
        columns={
            "s_info_name": "name",
            "s_ipo_listeddate": "list_date",
            "s_info_delistdate": "delist_date",
            "s_info_exchmarket": "exchange",
            "s_info_industry_sw": "industry_sw",
            "s_info_compname": "company_name",
            "s_info_province": "province",
            "s_info_business": "business",
        }
    )
    all_df = all_df.drop_duplicates("symbol").sort_values("symbol")
    output.parent.mkdir(parents=True, exist_ok=True)
    all_df.to_csv(output, index=False, encoding="utf-8-sig")
    return all_df


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--daily", type=Path, default=DEFAULT_DAILY)
    parser.add_argument("--output", type=Path, default=Path("universe.csv"))
    parser.add_argument("--batch", type=int, default=100)
    args = parser.parse_args()
    df = build(args.daily, args.output, args.batch)
    print(f"wrote {len(df)} symbols to {args.output.resolve()}")


if __name__ == "__main__":
    main()
